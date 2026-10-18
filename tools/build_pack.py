"""Build the paid Excel workbook (Offboard Kit Pro Pack) from data.js.

Usage: python3 tools/build_pack.py [output.xlsx]
"""
import json
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "dist" / "Offboard-Kit-Pro-Pack.xlsx"

data = json.loads(subprocess.check_output(
    ["node", "-e", f'eval(require("fs").readFileSync({json.dumps(str(ROOT / "data.js"))},"utf8")'
     ' + ";console.log(JSON.stringify({SYSTEMS, GENERAL}))")']))

ACCENT = "4F6BFF"
INK = "0F1424"
MUTED = "7A8299"
BAND = "F4F6FB"
WHITE = "FFFFFF"
thin = Side(style="thin", color="E3E7F0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
H_FONT = Font(name="Calibri", bold=True, color=WHITE, size=11)
H_FILL = PatternFill("solid", fgColor=ACCENT)
SEC_FILL = PatternFill("solid", fgColor="E8ECFF")
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top")


def owner_for(step):
    s = step.lower()
    if "hr" in s.split() or "with hr" in s or "policy" in s:
        return "HR"
    if "manager" in s and ("agree" in s or "confirm" in s):
        return "Manager"
    if "badge" in s or "alarm" in s or "keys" in s:
        return "Facilities"
    return "IT"


def title_block(ws, title, subtitle, width):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    ws["A1"] = title
    ws["A1"].font = Font(name="Calibri", bold=True, size=18, color=INK)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Calibri", size=10, color=MUTED)
    ws.row_dimensions[1].height = 30


def details_block(ws, row, col=1):
    for i, label in enumerate(["Employee", "Department", "Manager", "Date", "Ticket"]):
        r = row + i
        ws.cell(r, col, label).font = Font(bold=True, color=MUTED)
        ws.merge_cells(start_row=r, start_column=col + 1, end_row=r, end_column=col + 2)
        c = ws.cell(r, col + 1)
        c.border = Border(bottom=Side(style="thin", color="CFD5E3"))
    return row + 6


def header_row(ws, row, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row, i, h)
        c.font = H_FONT
        c.fill = H_FILL
        c.alignment = Alignment(vertical="center", horizontal="left" if i > 1 else "center")
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 22


def checklist_sheet(wb, mode):
    kind = "Offboarding" if mode == "offboarding" else "Onboarding"
    ws = wb.create_sheet(kind)
    ws.sheet_view.showGridLines = False
    headers = ["#", "Area", "Step", "Owner", "Status", "Done on", "Done by", "Notes"]
    widths = [5, 24, 62, 11, 13, 12, 14, 34]
    title_block(ws, f"{kind} checklist", "Delete the areas your company doesn't use. Status: To do / Done / N/A.", len(headers))
    row = details_block(ws, 4, col=2)
    # Progress line, filled in after the rows are known.
    progress_row = row
    row += 2
    header_row(ws, row, headers, widths)
    first = row + 1
    row += 1
    sections = [("Before you start", data["GENERAL"][mode])] + [(s["label"], s[mode]) for s in data["SYSTEMS"]]
    n = 0
    for name, steps in sections:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
        c = ws.cell(row, 1, name)
        c.font = Font(bold=True, color=INK)
        c.fill = SEC_FILL
        row += 1
        for step in steps:
            n += 1
            text = step.replace("{name}", "the employee")
            values = [n, name, text, owner_for(step), "To do", None, None, None]
            for col, v in enumerate(values, start=1):
                cell = ws.cell(row, col, v)
                cell.border = BORDER
                cell.alignment = CENTER if col in (1, 4, 5) else WRAP
                if n % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor=BAND)
            ws.cell(row, 6).number_format = "yyyy-mm-dd"
            row += 1
    last = row - 1

    status = DataValidation(type="list", formula1='"To do,Done,N/A"', allow_blank=False)
    ws.add_data_validation(status)
    status.add(f"E{first}:E{last}")
    rng = f"E{first}:E{last}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Done"'], fill=PatternFill("solid", fgColor="D3F5E6"), font=Font(color="0B7A50", bold=True)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"N/A"'], fill=PatternFill("solid", fgColor="EEEEEE"), font=Font(color="888888")))
    ws.conditional_formatting.add(f"C{first}:C{last}", FormulaRule(formula=[f'$E{first}="Done"'], font=Font(color="888888", strike=True)))

    total = f'(COUNTA({rng})-COUNTIF({rng},"N/A"))'
    ws.cell(progress_row, 2, "Progress").font = Font(bold=True, color=MUTED)
    p = ws.cell(progress_row, 3, f'=TEXT(IFERROR(COUNTIF({rng},"Done")/{total},0),"0%")&"   ·   "&COUNTIF({rng},"Done")&" of "&{total}&" steps done"')
    p.font = Font(bold=True, size=13, color=ACCENT)

    ws.freeze_panes = ws.cell(first, 1)
    ws.print_title_rows = f"{first - 1}:{first - 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.auto_filter.ref = f"A{first - 1}:H{last}"


def table_sheet(wb, name, title, subtitle, headers, widths, rows=25, preset=None):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    title_block(ws, title, subtitle, len(headers))
    start = details_block(ws, 4)
    header_row(ws, start, headers, widths)
    preset = preset or []
    for i in range(max(rows, len(preset))):
        r = start + 1 + i
        values = preset[i] if i < len(preset) else [None] * len(headers)
        for col in range(1, len(headers) + 1):
            c = ws.cell(r, col, values[col - 1] if col - 1 < len(values) else None)
            c.border = BORDER
            c.alignment = WRAP
            if i % 2:
                c.fill = PatternFill("solid", fgColor=BAND)
    ws.freeze_panes = ws.cell(start + 1, 1)
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
    return ws, start


def yes_no(ws, col, start, rows):
    dv = DataValidation(type="list", formula1='"Yes,No,N/A"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col}{start + 1}:{col}{start + rows}")


wb = Workbook()
guide = wb.active
guide.title = "Start here"
guide.sheet_view.showGridLines = False
guide.column_dimensions["A"].width = 4
guide.column_dimensions["B"].width = 100
guide["B2"] = "Offboard Kit · Pro Pack"
guide["B2"].font = Font(bold=True, size=22, color=ACCENT)
guide["B3"] = "IT onboarding and offboarding workbook"
guide["B3"].font = Font(size=12, color=MUTED)
lines = [
    ("How to use it", True),
    ("1. Save a copy of this workbook for each person (e.g. Offboarding - Jane Smith - 2026-10-18).", False),
    ("2. Fill in the details at the top of the sheet you need.", False),
    ("3. Delete the areas your company doesn't use, or mark their steps N/A.", False),
    ("4. Set each step's Status to Done as you go. Progress updates automatically.", False),
    ("5. Keep the finished copy with the ticket as your audit record.", False),
    ("", False),
    ("What's inside", True),
    ("Offboarding: every step to cut access, hand over data and recover equipment, in the right order.", False),
    ("Onboarding: accounts, groups, devices and access for a new starter.", False),
    ("Exit-day timeline: what to do before, on and after the last day.", False),
    ("Asset return: a form to record every device and item handed back.", False),
    ("Access review: a log of every account and when it was disabled, for audits.", False),
    ("Shared credentials: the shared passwords to rotate after someone leaves.", False),
    ("", False),
    ("Good to know", True),
    ("Cut access first (sign-in, sessions, VPN), then handle data, then hardware and licenses.", False),
    ("Check every step against your company's own policies and data retention rules.", False),
    ("This workbook stores nothing online. Keep completed copies according to your company's records policy.", False),
]
for i, (text, bold) in enumerate(lines, start=5):
    c = guide.cell(i, 2, text)
    c.font = Font(bold=bold, size=12 if bold else 11, color=INK if bold else "3A4258")
    c.alignment = Alignment(wrap_text=True)

checklist_sheet(wb, "offboarding")
checklist_sheet(wb, "onboarding")

timeline = [
    ["Before last day", "Confirm last day and cut-off time with HR", "HR / IT"],
    ["Before last day", "Agree with the manager where files and email go", "Manager"],
    ["Before last day", "List every system and shared credential the person uses", "IT"],
    ["Before last day", "Schedule device collection", "IT"],
    ["Last day", "Collect laptop, phone, badge and peripherals", "IT / Facilities"],
    ["Last day", "Block sign-in and revoke sessions at the agreed time", "IT"],
    ["Last day", "Disable VPN, remote access and building access", "IT / Facilities"],
    ["Last day", "Set out-of-office and mail forwarding", "IT"],
    ["After", "Transfer files and mailbox to the manager", "IT"],
    ["After", "Rotate shared passwords", "IT"],
    ["After", "Wipe devices and update asset inventory", "IT"],
    ["After", "Remove licenses and paid seats", "IT"],
    ["After", "Delete or archive accounts per retention policy (30-90 days)", "IT"],
]
ws, s = table_sheet(wb, "Exit-day timeline", "Exit-day timeline", "Before, on and after the last day.",
                    ["When", "Task", "Owner", "Done", "Notes"], [18, 60, 16, 10, 34], rows=len(timeline), preset=timeline)
yes_no(ws, "D", s, len(timeline))

ws, s = table_sheet(wb, "Asset return", "Asset return form", "Record every item handed back.",
                    ["Item", "Make / model", "Serial / asset tag", "Condition", "Returned", "Date", "Received by"],
                    [22, 24, 22, 16, 11, 12, 18], rows=15,
                    preset=[["Laptop"], ["Charger"], ["Monitor"], ["Dock"], ["Headset"], ["Phone"], ["Access badge"], ["Keys"]])
yes_no(ws, "E", s, 15)

ws, s = table_sheet(wb, "Access review", "Access review log", "Every account the person had and when it was disabled.",
                    ["System / app", "Username", "Access level", "Disabled", "Date", "By", "Data transferred to"],
                    [24, 22, 18, 11, 12, 14, 24], rows=25,
                    preset=[["Microsoft 365 / Google Workspace"], ["Active Directory"], ["VPN"], ["Slack / Teams"], ["Password manager"], ["CRM"], ["Accounting"]])
yes_no(ws, "D", s, 25)

ws, s = table_sheet(wb, "Shared credentials", "Shared credentials to rotate", "Shared logins the person knew. Never write the passwords themselves here.",
                    ["Credential", "Where it's stored", "Rotated", "Date", "By"], [34, 30, 11, 12, 16], rows=20,
                    preset=[["Office Wi-Fi"], ["Guest Wi-Fi"], ["Firewall / router admin"], ["Domain registrar"], ["Vendor portals"], ["Alarm code"]])
yes_no(ws, "C", s, 20)

for sheet in wb.worksheets:
    sheet.sheet_properties.tabColor = ACCENT if sheet.title in ("Offboarding", "Onboarding") else "CFD5E3"

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(OUT)
